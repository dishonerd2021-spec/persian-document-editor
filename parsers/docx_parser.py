"""
DOCX parser for Word documents
"""

from pathlib import Path
from typing import Optional
from loguru import logger

try:
    from docx import Document as DocxDocument
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    DocxDocument = None
    logger.warning("python-docx not installed")

from core.document import Document
from core.paragraph import Paragraph, ParagraphStyle, TextDirection


class DocxParser:
    """Parser for DOCX (Word) files"""
    
    def __init__(self):
        """Initialize DOCX parser"""
        if DocxDocument is None:
            logger.error("python-docx is required for DOCX parsing")
    
    @staticmethod
    def parse_docx(file_path: str, title: Optional[str] = None) -> Document:
        """
        Parse DOCX file into Document
        
        Args:
            file_path: Path to DOCX file
            title: Document title (uses filename if not specified)
            
        Returns:
            Document object
        """
        if DocxDocument is None:
            logger.error("python-docx not installed")
            return Document()
        
        try:
            file_path = Path(file_path)
            
            if not file_path.exists():
                logger.error(f"File not found: {file_path}")
                return Document()
            
            docx_doc = DocxDocument(str(file_path))
            
            if title is None:
                title = file_path.stem
            
            doc = Document(title=title)
            
            # Extract text from paragraphs
            for para in docx_doc.paragraphs:
                if not para.text.strip():
                    continue
                
                # Determine style
                style = ParagraphStyle.NORMAL
                if para.style and 'Title' in para.style.name:
                    style = ParagraphStyle.TITLE
                elif para.style and 'Heading' in para.style.name:
                    if 'Heading 1' in para.style.name:
                        style = ParagraphStyle.HEADING1
                    elif 'Heading 2' in para.style.name:
                        style = ParagraphStyle.HEADING2
                    else:
                        style = ParagraphStyle.HEADING3
                
                # Create paragraph
                new_para = doc.add_paragraph(para.text.strip(), style)
                
                # Copy runs for better formatting
                for run in para.runs:
                    if run.text.strip():
                        new_para.add_text(
                            run.text,
                            bold=run.bold,
                            italic=run.italic,
                            font_size=run.font.size.pt if run.font.size else None,
                            font_name=run.font.name
                        )
            
            # Extract tables
            for table in docx_doc.tables:
                table_data = []
                for row in table.rows:
                    row_data = [cell.text for cell in row.cells]
                    table_data.append(row_data)
                
                if table_data:
                    doc.add_table(table_data, "Table")
            
            logger.info(f"Parsed DOCX file: {file_path}")
            return doc
        except Exception as e:
            logger.error(f"Error parsing DOCX file: {e}")
            return Document(title=title or "Untitled")
    
    @staticmethod
    def export_to_docx(doc: Document, output_path: str) -> bool:
        """
        Export Document to DOCX format
        
        Args:
            doc: Document to export
            output_path: Path to save DOCX file
            
        Returns:
            True if successful
        """
        if DocxDocument is None:
            logger.error("python-docx not installed")
            return False
        
        try:
            docx_doc = DocxDocument()
            
            # Add title
            if doc.metadata.title:
                title_para = docx_doc.add_paragraph(doc.metadata.title)
                title_para.style = 'Title'
            
            # Add abstract
            if doc.metadata.abstract:
                abs_heading = docx_doc.add_paragraph("Abstract")
                abs_heading.style = 'Heading 1'
                docx_doc.add_paragraph(doc.metadata.abstract)
            
            # Add keywords
            if doc.metadata.keywords:
                kw_heading = docx_doc.add_paragraph("Keywords")
                kw_heading.style = 'Heading 1'
                kw_text = ", ".join(doc.metadata.keywords)
                docx_doc.add_paragraph(kw_text)
            
            # Add paragraphs
            for para in doc.paragraphs:
                if para.is_empty():
                    continue
                
                # Map style
                style_map = {
                    ParagraphStyle.TITLE: 'Title',
                    ParagraphStyle.HEADING1: 'Heading 1',
                    ParagraphStyle.HEADING2: 'Heading 2',
                    ParagraphStyle.HEADING3: 'Heading 3',
                }
                
                docx_para = docx_doc.add_paragraph(para.get_text())
                if para.format.style in style_map:
                    docx_para.style = style_map[para.format.style]
            
            # Add tables
            for table in doc.tables:
                docx_table = docx_doc.add_table(rows=len(table['data']), cols=len(table['data'][0]) if table['data'] else 0)
                for i, row in enumerate(table['data']):
                    for j, cell_text in enumerate(row):
                        docx_table.rows[i].cells[j].text = str(cell_text)
            
            # Add references
            if doc.references:
                ref_heading = docx_doc.add_paragraph("References")
                ref_heading.style = 'Heading 1'
                for ref in doc.references:
                    ref_text = ref.get('title', '') or ref.get('author', '')
                    docx_doc.add_paragraph(ref_text)
            
            # Save document
            output_path = Path(output_path)
            output_path.parent.mkdir(parents=True, exist_ok=True)
            docx_doc.save(str(output_path))
            
            logger.info(f"Exported document to DOCX: {output_path}")
            return True
        except Exception as e:
            logger.error(f"Error exporting to DOCX: {e}")
            return False
