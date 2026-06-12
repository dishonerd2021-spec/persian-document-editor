"""
DOCX exporter for Word document format
"""

from typing import Optional, Dict, Any
from pathlib import Path
from loguru import logger

from core.document import Document
from core.paragraph import ParagraphStyle, TextDirection
from exporters.base_exporter import BaseExporter, EXPORT_FORMATS, FONT_CONFIG

try:
    from docx import Document as DocxDocument
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    DocxDocument = None
    logger.warning("python-docx not installed")


class DocxExporter(BaseExporter):
    """Exporter for DOCX (Word) format"""
    
    def __init__(self):
        """Initialize DOCX exporter"""
        super().__init__()
        self.output_format = "docx"
        
        if DocxDocument is None:
            logger.error("python-docx is required for DOCX export")
    
    def export(
        self,
        document: Document,
        output_path: str,
        format_name: str = 'ieee',
        **kwargs
    ) -> bool:
        """
        Export document to DOCX format
        
        Args:
            document: Document to export
            output_path: Path to save DOCX file
            format_name: Export format (ieee, acm, springer, elsevier, thesis)
            **kwargs: Additional options
            
        Returns:
            True if successful
        """
        if DocxDocument is None:
            logger.error("python-docx is not installed")
            return False
        
        try:
            # Ensure output directory exists
            if not self.ensure_output_dir(output_path):
                return False
            
            # Create DOCX document
            docx_doc = DocxDocument()
            
            # Get format settings
            format_settings = EXPORT_FORMATS.get(format_name, EXPORT_FORMATS['ieee'])
            
            # Set up margins
            sections = docx_doc.sections
            for section in sections:
                section.top_margin = Inches(format_settings['margin_top'] / 10)
                section.bottom_margin = Inches(format_settings['margin_bottom'] / 10)
                section.left_margin = Inches(format_settings['margin_left'] / 10)
                section.right_margin = Inches(format_settings['margin_right'] / 10)
            
            # Add title
            if document.metadata.title:
                title_para = docx_doc.add_paragraph(document.metadata.title)
                title_para.style = 'Title'
                title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            # Add author
            if document.metadata.author:
                author_para = docx_doc.add_paragraph(document.metadata.author)
                author_para.style = 'Subtitle'
                author_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            # Add abstract
            if document.metadata.abstract:
                abstract_heading = docx_doc.add_paragraph()
                abstract_heading.add_run("Abstract").bold = True
                abstract_para = docx_doc.add_paragraph(document.metadata.abstract)
                abstract_para.paragraph_format.left_indent = Inches(0.25)
            
            # Add keywords
            if document.metadata.keywords:
                keywords_heading = docx_doc.add_paragraph()
                keywords_heading.add_run("Keywords: ").bold = True
                keywords_text = ", ".join(document.metadata.keywords)
                keywords_para = docx_doc.add_paragraph(keywords_text)
                keywords_para.paragraph_format.left_indent = Inches(0.25)
            
            # Add content paragraphs
            for para in document.paragraphs:
                if para.is_empty():
                    continue
                
                # Create paragraph
                docx_para = docx_doc.add_paragraph(para.get_text())
                
                # Apply style
                style_map = {
                    ParagraphStyle.TITLE: 'Title',
                    ParagraphStyle.HEADING1: 'Heading 1',
                    ParagraphStyle.HEADING2: 'Heading 2',
                    ParagraphStyle.HEADING3: 'Heading 3',
                    ParagraphStyle.ABSTRACT: 'Normal',
                }
                
                if para.format.style in style_map:
                    docx_para.style = style_map[para.format.style]
                
                # Apply alignment based on direction
                if para.format.direction == TextDirection.RTL:
                    docx_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                elif para.format.direction == TextDirection.LTR:
                    docx_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
                else:
                    docx_para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                
                # Apply formatting
                if para.format.line_spacing:
                    docx_para.paragraph_format.line_spacing = para.format.line_spacing
                if para.format.space_after:
                    docx_para.paragraph_format.space_after = Pt(para.format.space_after)
                if para.format.space_before:
                    docx_para.paragraph_format.space_before = Pt(para.format.space_before)
            
            # Add tables
            for table_data in document.tables:
                if 'data' not in table_data or not table_data['data']:
                    continue
                
                data = table_data['data']
                rows = len(data)
                cols = len(data[0]) if data else 0
                
                docx_table = docx_doc.add_table(rows=rows, cols=cols)
                docx_table.style = 'Light Grid Accent 1'
                
                for i, row in enumerate(data):
                    for j, cell_text in enumerate(row):
                        docx_table.rows[i].cells[j].text = str(cell_text)
                
                # Add caption
                if table_data.get('caption'):
                    caption_para = docx_doc.add_paragraph()
                    caption_run = caption_para.add_run(table_data['caption'])
                    caption_run.italic = True
                    caption_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            # Add references
            if document.references:
                docx_doc.add_paragraph()  # Empty line
                ref_heading = docx_doc.add_paragraph()
                ref_heading.add_run("References").bold = True
                
                for i, ref in enumerate(document.references, 1):
                    ref_text = ref.get('title', '') or ref.get('author', '')
                    ref_para = docx_doc.add_paragraph(f"[{i}] {ref_text}")
                    ref_para.paragraph_format.left_indent = Inches(0.25)
                    ref_para.paragraph_format.hanging_indent = Inches(-0.25)
            
            # Save document
            output_path = Path(output_path)
            if not str(output_path).endswith('.docx'):
                output_path = output_path.with_suffix('.docx')
            
            docx_doc.save(str(output_path))
            logger.info(f"Document exported to DOCX: {output_path}")
            return True
        except Exception as e:
            logger.error(f"Error exporting to DOCX: {e}")
            return False
    
    def export_batch(
        self,
        documents: list,
        output_dir: str,
        format_name: str = 'ieee'
    ) -> Dict[str, bool]:
        """
        Export multiple documents
        
        Args:
            documents: List of documents to export
            output_dir: Directory to save files
            format_name: Export format
            
        Returns:
            Dictionary with export results
        """
        results = {}
        
        for doc in documents:
            output_path = Path(output_dir) / f"{doc.metadata.title or 'document'}.docx"
            success = self.export(doc, str(output_path), format_name)
            results[doc.metadata.title or 'unknown'] = success
        
        return results
